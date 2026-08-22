"""
Master Excel Register Log Exporter using openpyxl.
"""

import datetime
import os
import tempfile
import threading
from contextlib import contextmanager
from typing import Any, Dict, Iterator, List, Optional

from openpyxl import Workbook, load_workbook

from dp_lookup.exporters.contracts import ExcelRegisterData

try:  # pragma: no cover - platform-specific import
    import fcntl
except ImportError:  # pragma: no cover - Windows
    fcntl = None  # type: ignore[assignment]

try:  # pragma: no cover - platform-specific import
    import msvcrt
except ImportError:  # pragma: no cover - Unix
    msvcrt = None  # type: ignore[assignment]


_REGISTER_LOCK = threading.RLock()


@contextmanager
def _register_file_lock(register_path: str) -> Iterator[None]:
    """Hold a stable advisory lock beside the register while it is rewritten."""
    lock_path = f"{register_path}.lock"
    lock_handle = None
    try:
        lock_handle = open(lock_path, "a+b")
        if fcntl is not None:
            try:
                fcntl.flock(lock_handle.fileno(), fcntl.LOCK_EX)
            except OSError:
                pass
        elif msvcrt is not None:
            # msvcrt.locking locks bytes, so ensure the lock file has one.
            lock_handle.seek(0)
            lock_handle.write(b"0")
            lock_handle.flush()
            lock_handle.seek(0)
            try:
                msvcrt.locking(lock_handle.fileno(), msvcrt.LK_LOCK, 1)
            except OSError:
                pass
    except OSError:
        # The process-local lock still protects threads if an advisory lock is
        # unavailable on a particular filesystem.
        pass
    try:
        yield
    finally:
        if lock_handle is not None:
            if fcntl is not None:
                try:
                    fcntl.flock(lock_handle.fileno(), fcntl.LOCK_UN)
                except OSError:
                    pass
            elif msvcrt is not None:
                try:
                    lock_handle.seek(0)
                    msvcrt.locking(lock_handle.fileno(), msvcrt.LK_UNLCK, 1)
                except OSError:
                    pass
            try:
                lock_handle.close()
            except OSError:
                pass


def _save_workbook_atomic(workbook: Workbook, register_path: str) -> None:
    """Save a workbook beside the target, then atomically replace the target."""
    directory = os.path.dirname(register_path) or "."
    descriptor, temporary_path = tempfile.mkstemp(
        prefix=f".{os.path.basename(register_path)}.",
        suffix=".tmp",
        dir=directory,
    )
    os.close(descriptor)
    try:
        workbook.save(temporary_path)
        os.replace(temporary_path, register_path)
    finally:
        try:
            os.unlink(temporary_path)
        except OSError:
            pass


def update_excel_register(
    register_path: str,
    attrs: Dict[str, Any],
    village: str,
    cts_number: str,
    zone: str,
    status_badge: str,
    res_code: str,
    res_type: str,
    des_code: str,
    des_desc: str,
    mod_approval: str,
    mod_label: str,
    mod_doc: str,
    crz_buffer_flag: str,
    metro_buffer_flag: str,
    road_name: str,
    road_width: str,
    neighbors: List[Dict[str, Any]],
    dp_snapshot_fname: str,
    sat_snapshot_fname: str,
    pdf_fname: str,
    dxf_fname: str,
    geojson_fname: str,
    kml_fname: str,
    map_link: str,
    reservations: Optional[List[Dict[str, Any]]] = None,
    designations: Optional[List[Dict[str, Any]]] = None,
) -> None:
    """Append a structured lookup record row to the master Excel register."""
    with _REGISTER_LOCK:
        now_str = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")

        headers = [
            "lookup_datetime",
            "source",
            "ward",
            "village",
            "cts_no",
            "type",
            "area_sqm",
            "zone",
            "status_badge",
            "reservation_code",
            "reservation_type",
            "designation_code",
            "designation_desc",
            "modification_approval",
            "modification_label",
            "modification_doc",
            "crz_buffer",
            "metro_buffer",
            "abutting_road",
            "road_width",
            "adjoining_plots_count",
            "hd_snapshot_file",
            "satellite_snapshot_file",
            "pdf_report_file",
            "dxf_file",
            "geojson_file",
            "kml_file",
            "map_link",
            "recorded_plot_area_sqm",
            "calculated_boundary_area_sqm",
            "area_difference_sqm",
            "area_difference_percent",
            "area_reliability",
            "area_quality_flags",
            "area_geometry_source",
            "area_source_snapshot_id",
            "area_source_snapshot_date",
            "authority_status",
            "authority_policy_version",
            "authority_area_record_id",
            "authority_geometry_record_id",
            "authority_planning_record_ids",
            "authority_active_record_ids",
        ]

        assessment = attrs.get("_DP_AREA_ASSESSMENT")
        if not isinstance(assessment, dict):
            assessment = {}
        selected_area = assessment.get("selected_area_sqm", attrs.get("AREA_APP_SQ_MTRS"))
        quality_flags = assessment.get("quality_flags") or []
        quality_flag_codes = "; ".join(
            str(flag.get("code")) for flag in quality_flags if isinstance(flag, dict) and flag.get("code")
        )
        authority = attrs.get("_DP_AUTHORITY_RESOLUTION")
        if not isinstance(authority, dict):
            authority = {}
        raw_area_authority = authority.get("area")
        area_authority = raw_area_authority if isinstance(raw_area_authority, dict) else {}
        raw_geometry_authority = authority.get("geometry")
        geometry_authority = raw_geometry_authority if isinstance(raw_geometry_authority, dict) else {}
        raw_planning_authority = authority.get("planning")
        planning_authority = raw_planning_authority if isinstance(raw_planning_authority, list) else []
        planning_record_ids = "; ".join(
            str(item.get("record_id"))
            for item in planning_authority
            if isinstance(item, dict) and item.get("record_id") is not None
        )
        raw_active_record_ids = authority.get("active_record_ids")
        active_record_ids = (
            "; ".join(str(value) for value in raw_active_record_ids)
            if isinstance(raw_active_record_ids, list)
            else ""
        )

        excel_res_code = "; ".join(r["code"] for r in reservations) if reservations else res_code
        excel_res_type = (
            "; ".join(f"{r['type']} ({r['description']})" for r in reservations) if reservations else res_type
        )
        excel_des_code = "; ".join(d["code"] for d in designations) if designations else des_code
        excel_des_desc = "; ".join(d["description"] for d in designations) if designations else des_desc

        row = [
            now_str,
            "MCGM SDP 2014-34",
            attrs.get("WARD"),
            village,
            cts_number,
            attrs.get("TYPE"),
            selected_area,
            zone,
            status_badge,
            excel_res_code,
            excel_res_type,
            excel_des_code,
            excel_des_desc,
            mod_approval,
            mod_label,
            mod_doc,
            crz_buffer_flag,
            metro_buffer_flag,
            road_name,
            road_width,
            len(neighbors),
            dp_snapshot_fname,
            sat_snapshot_fname,
            pdf_fname,
            dxf_fname,
            geojson_fname,
            kml_fname,
            map_link,
            assessment.get("recorded_plot_area_sqm", attrs.get("AREA_APP_SQ_MTRS")),
            assessment.get("calculated_boundary_area_sqm"),
            assessment.get("difference_sqm"),
            assessment.get("difference_percent"),
            assessment.get("reliability_label", "Screening"),
            quality_flag_codes,
            assessment.get("geometry_source", "MCGM cadastral parcel boundary"),
            assessment.get("source_snapshot_id"),
            assessment.get("source_snapshot_date"),
            authority.get("status"),
            authority.get("policy_version"),
            area_authority.get("record_id"),
            geometry_authority.get("record_id"),
            planning_record_ids,
            active_record_ids,
        ]

        register_path = os.path.realpath(os.path.abspath(os.path.expanduser(os.fspath(register_path))))
        os.makedirs(os.path.dirname(register_path), exist_ok=True)
        with _register_file_lock(register_path):
            wb = load_workbook(register_path) if os.path.exists(register_path) else Workbook()
            try:
                ws = wb.active
                if ws.max_row == 1 and ws["A1"].value is None:
                    for column, header in enumerate(headers, start=1):
                        ws.cell(row=1, column=column, value=header)
                else:
                    existing_headers = {
                        str(ws.cell(row=1, column=column).value): column
                        for column in range(1, ws.max_column + 1)
                        if ws.cell(row=1, column=column).value is not None
                    }
                    for column, header in enumerate(headers, start=1):
                        if header not in existing_headers:
                            ws.cell(row=1, column=column, value=header)
                ws.append(row)
                _save_workbook_atomic(wb, register_path)
            finally:
                wb.close()


def write_excel_register(data: ExcelRegisterData) -> None:
    """Append one row from the pipeline's typed report contract."""
    report = data.report
    names = data.artifact_filenames
    update_excel_register(
        register_path=report.output_path,
        attrs=dict(report.attrs),
        village=report.village,
        cts_number=report.cts_number,
        zone=report.zone,
        status_badge=report.status_badge,
        res_code=report.reservation_code,
        res_type=report.reservation_type,
        des_code=report.designation_code,
        des_desc=report.designation_description,
        mod_approval=report.modification_approval,
        mod_label=report.modification_label,
        mod_doc=report.modification_document,
        crz_buffer_flag=report.crz_status,
        metro_buffer_flag=report.metro_buffer,
        road_name=report.road_name,
        road_width=report.road_width,
        neighbors=[dict(item) for item in report.neighbours],
        dp_snapshot_fname=str(names["dp_map"]),
        sat_snapshot_fname=str(names["satellite"]),
        pdf_fname=str(names["pdf"]),
        dxf_fname=str(names["dxf"]),
        geojson_fname=str(names["geojson"]),
        kml_fname=str(names["kml"]),
        map_link=report.map_link,
        reservations=[dict(item) for item in report.reservations],
        designations=[dict(item) for item in report.designations],
    )
