"""Process-safe and atomic master-register persistence tests."""

import multiprocessing
from pathlib import Path

import openpyxl
import pytest

from dp_lookup.exporters import excel


def _register_args(path, index):
    return {
        "register_path": str(path),
        "attrs": {"WARD": "G/S", "TYPE": "CTS", "AREA_APP_SQ_MTRS": 100.0 + index},
        "village": "WORLI",
        "cts_number": str(700 + index),
        "zone": "R",
        "status_badge": "CLEAR",
        "res_code": "None",
        "res_type": "None",
        "des_code": "None",
        "des_desc": "None",
        "mod_approval": "None",
        "mod_label": "None",
        "mod_doc": "None",
        "crz_buffer_flag": "NO",
        "metro_buffer_flag": "NO",
        "road_name": "Road",
        "road_width": "12 M",
        "neighbors": [],
        "dp_snapshot_fname": "dp.png",
        "sat_snapshot_fname": "sat.png",
        "pdf_fname": "report.pdf",
        "dxf_fname": "plot.dxf",
        "geojson_fname": "plot.geojson",
        "kml_fname": "plot.kml",
        "map_link": "https://example.invalid/plot",
    }


def _child_append(args):
    path, index = args
    from dp_lookup.exporters.excel import update_excel_register

    update_excel_register(**_register_args(path, index))


def test_register_preserves_existing_xlsx_shape(tmp_path):
    path = tmp_path / "dp-lookups.xlsx"
    excel.update_excel_register(**_register_args(path, 0))

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == 2
    assert sheet.max_column == 43
    assert sheet.cell(row=1, column=29).value == "recorded_plot_area_sqm"
    assert sheet.cell(row=1, column=30).value == "calculated_boundary_area_sqm"
    assert sheet.cell(1, 5).value == "cts_no"
    assert sheet.cell(2, 5).value == "700"
    workbook.close()


def test_failed_atomic_save_preserves_previous_register(tmp_path, monkeypatch):
    path = tmp_path / "dp-lookups.xlsx"
    excel.update_excel_register(**_register_args(path, 0))
    before = path.read_bytes()

    def fail_replace(_source, _target):
        raise OSError("simulated replace failure")

    monkeypatch.setattr(excel.os, "replace", fail_replace)
    with pytest.raises(OSError, match="simulated replace failure"):
        excel.update_excel_register(**_register_args(path, 1))

    assert path.read_bytes() == before
    assert not list(tmp_path.glob("*.tmp"))


def test_register_appends_structured_area_authority_columns(tmp_path):
    path = tmp_path / "dp-lookups.xlsx"
    args = _register_args(path, 0)
    args["attrs"]["_DP_AREA_ASSESSMENT"] = {
        "recorded_plot_area_sqm": 100.0,
        "calculated_boundary_area_sqm": 106.0,
        "selected_area_sqm": 100.0,
        "difference_sqm": 6.0,
        "difference_percent": 6.0,
        "reliability_label": "Qualified with warnings",
        "geometry_source": "MCGM cadastral parcel boundary",
        "quality_flags": [{"code": "area_difference_over_5_percent"}],
        "source_snapshot_id": 7,
        "source_snapshot_date": "2026-08-22T00:00:00+00:00",
    }

    excel.update_excel_register(**args)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet = workbook.active
    values = {sheet.cell(1, column).value: sheet.cell(2, column).value for column in range(1, 44)}
    workbook.close()
    assert values["recorded_plot_area_sqm"] == 100.0
    assert values["calculated_boundary_area_sqm"] == 106.0
    assert values["area_difference_percent"] == 6.0
    assert values["area_reliability"] == "Qualified with warnings"
    assert values["area_quality_flags"] == "area_difference_over_5_percent"
    assert values["area_source_snapshot_id"] == 7


def test_multiprocess_register_appends_do_not_lose_rows(tmp_path):
    path = str(tmp_path / "dp-lookups.xlsx")
    jobs = [(path, index) for index in range(8)]
    context = multiprocessing.get_context("spawn")
    with context.Pool(processes=4) as pool:
        pool.map(_child_append, jobs)

    workbook = openpyxl.load_workbook(path, read_only=True, data_only=False)
    sheet = workbook.active
    assert sheet.max_row == 9
    assert {sheet.cell(row, 5).value for row in range(2, 10)} == {str(700 + i) for i in range(8)}
    workbook.close()
    assert Path(f"{path}.lock").exists()
